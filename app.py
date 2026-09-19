import base64
import io
import ipaddress
import json
import logging
import os
import re
import socket
import xml.etree.ElementTree as ET
from collections import Counter
from copy import deepcopy
from urllib.parse import parse_qs, quote_plus, unquote, urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from flask import Flask, jsonify, render_template, request, send_file
from pypdf import PdfReader
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.util import Inches, Pt
from ppt_layout import fit_text, readable_color, text_background, style_frame


app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024
app.logger.setLevel(logging.INFO)


STRUCTURE_LIBRARY = [
    {
        "id": "odi_accessibility_match",
        "name": "Accessibility-first capability match",
        "source_pair": "ODI Visual Designer",
        "best_for": "Specialist or individual-contractor RFPs",
        "section_order": ["problem", "qualification match", "evidence", "commercials"],
        "phases": [
            {"name": "Align", "duration": "1 week", "goal": "Confirm users, accessibility obligations, and success measures."},
            {"name": "Design", "duration": "2-3 weeks", "goal": "Translate research and requirements into feasible, inclusive design."},
            {"name": "Deliver", "duration": "1-2 weeks", "goal": "Hand over production-ready assets and reusable standards."},
        ],
        "evidence": ["8+ years of directly relevant experience", "Public-sector-adjacent delivery", "Accessibility and inclusive-design track record"],
        "risk_note": "Validate accessibility and stakeholder acceptance continuously rather than at final review.",
        "artifact": "Reusable design standards and production-ready assets",
    },
    {
        "id": "nj_assess_design_implement",
        "name": "Assess → Design → Implement",
        "source_pair": "McKinsey / New Jersey DMAVA",
        "best_for": "Organization redesign and politically visible transformation",
        "section_order": ["problem", "assessment", "target design", "implementation", "why us"],
        "phases": [
            {"name": "Assess", "duration": "6 weeks", "goal": "Map the current operating model and benchmark relevant precedents."},
            {"name": "Design", "duration": "10 weeks", "goal": "Co-design the future model, roles, governance, and economics."},
            {"name": "Implement", "duration": "8 weeks", "goal": "Build a costed transition roadmap with milestones and owners."},
        ],
        "evidence": ["5,460 public-sector engagements", "Experience across 105+ countries", "Named government reorganization precedent"],
        "risk_note": "Protect service continuity while the future operating model is designed and transitioned.",
        "artifact": "Costed target operating model and milestone-based implementation roadmap",
    },
    {
        "id": "wa_learn_grow_apply",
        "name": "Learn → Grow → Apply",
        "source_pair": "McKinsey / Washington State",
        "best_for": "Capability-building, roster, and change-management RFPs",
        "section_order": ["need", "capability method", "quality assurance", "capacity transfer"],
        "phases": [
            {"name": "Learn", "duration": "2 weeks", "goal": "Build a shared fact base and introduce practical methods."},
            {"name": "Grow", "duration": "3 weeks", "goal": "Coach teams through feedback, iteration, and quality review."},
            {"name": "Apply", "duration": "4 weeks", "goal": "Use the methods on live work and transfer ownership to client teams."},
        ],
        "evidence": ["Experience with 25 U.S. states", "Two-partner quality review", "Explicit capacity-transfer methodology"],
        "risk_note": "Measure success by client self-sufficiency, not ongoing consultant dependency.",
        "artifact": "Client-owned playbooks, training materials, and quality-assurance process",
    },
    {
        "id": "macc_model_and_readiness",
        "name": "Benchmark → Model → Ready → Refine",
        "source_pair": "BCG / Massachusetts Community Colleges",
        "best_for": "Policy design, scenario modeling, and implementation readiness",
        "section_order": ["problem", "options", "impact model", "readiness", "stakeholder refinement"],
        "phases": [
            {"name": "Benchmark", "duration": "2 weeks", "goal": "Compare relevant programs and define viable design options."},
            {"name": "Model", "duration": "3 weeks", "goal": "Build baseline and scenario economics with explicit trade-offs."},
            {"name": "Ready", "duration": "3 weeks", "goal": "Assess operational capacity, processes, and resource gaps."},
            {"name": "Refine", "duration": "4-8 weeks", "goal": "Refine the preferred option with stakeholders using an editable model."},
        ],
        "evidence": ["100+ higher-education institutions", "Named practitioners with implementation experience", "Demonstrated rapid mobilization"],
        "risk_note": "Keep assumptions editable so stakeholders can test trade-offs before committing.",
        "artifact": "Editable scenario model and implementation-readiness assessment",
    },
    {
        "id": "berkeley_fact_base",
        "name": "Fact base → Opportunities → Options → Vet",
        "source_pair": "Bain / UC Berkeley",
        "best_for": "Cost transformation with complex governance and explicit scope limits",
        "section_order": ["stakes", "scope fence", "fact base", "opportunities", "business cases", "risk"],
        "phases": [
            {"name": "Build the fact base", "duration": "5-6 weeks", "goal": "Create a trusted baseline, operating metrics, and reusable data asset."},
            {"name": "Identify opportunities", "duration": "6-8 weeks", "goal": "Triangulate analysis, benchmarks, and stakeholder workshops."},
            {"name": "Design options", "duration": "8-10 weeks", "goal": "Build costed options with impact, timing, barriers, and owners."},
            {"name": "Vet options", "duration": "2-4 weeks", "goal": "Test options with stakeholders before final prioritization."},
            {"name": "Manage change", "duration": "Continuous", "goal": "Maintain alignment, communications, and governance throughout."},
        ],
        "evidence": ["Purpose-built cost-cube method", "Triangulated benchmarking", "Realistic 50-70% realization assumption"],
        "risk_note": "Do not promise 100% realization; agree a realistic yield and protect explicit out-of-scope areas.",
        "artifact": "Reusable cost-cube database and prioritized portfolio of business-cased options",
    },
    {
        "id": "virginia_spend_transformation",
        "name": "Baseline → Playbooks → Execute",
        "source_pair": "BCG / Virginia Procurement",
        "best_for": "Procurement, spend, and rapid savings transformation",
        "section_order": ["quantified opportunity", "clean baseline", "category playbooks", "execution", "risk sharing"],
        "phases": [
            {"name": "Build the baseline", "duration": "90 days", "goal": "Clean and reconcile multi-year data into a trusted spend baseline."},
            {"name": "Build playbooks", "duration": "60 days", "goal": "Prioritize opportunities and create agency-level category playbooks."},
            {"name": "Execute", "duration": "60 days", "goal": "Launch quick wins and transfer repeatable savings capabilities."},
        ],
        "evidence": ["990 procurement projects in five years", "18 state-government clients", "25 pre-built category playbooks"],
        "risk_note": "Tie ambition to validated baseline data and stage-gate savings claims before execution.",
        "artifact": "Clean spend cube, quantified opportunity pipeline, and reusable category playbooks",
    },
]

NEUTRAL_BRAND = {
    "primary": "#172033",
    "secondary": "#F5F2EA",
    "accent": "#D46A3A",
    "logo_url": "",
    "logo_data_url": "",
    "tone": "Clear, credible, and evidence-led",
    "source": "neutral fallback",
}

STATE_NAMES = [
    "Alabama", "Alaska", "Arizona", "Arkansas", "California", "Colorado", "Connecticut",
    "Delaware", "Florida", "Georgia", "Hawaii", "Idaho", "Illinois", "Indiana", "Iowa",
    "Kansas", "Kentucky", "Louisiana", "Maine", "Maryland", "Massachusetts", "Michigan",
    "Minnesota", "Mississippi", "Missouri", "Montana", "Nebraska", "Nevada", "New Hampshire",
    "New Jersey", "New Mexico", "New York", "North Carolina", "North Dakota", "Ohio",
    "Oklahoma", "Oregon", "Pennsylvania", "Rhode Island", "South Carolina", "South Dakota",
    "Tennessee", "Texas", "Utah", "Vermont", "Virginia", "Washington", "West Virginia",
    "Wisconsin", "Wyoming",
]

KNOWN_BRAND_OWNERS = [
    {
        "owner": "ITC",
        "sector": "FMCG",
        "brands": ["ashirvaad atta", "aashirvaad atta", "ashirvaad", "aashirvaad", "sunfeast", "classmate", "fiama"],
    },
    {
        "owner": "Tata Consumer Products",
        "sector": "FMCG",
        "brands": ["tata salt", "tata tea", "tetley", "himalayan water"],
    },
    {
        "owner": "Britannia Industries",
        "sector": "FMCG",
        "brands": ["britannia", "good day", "tiger biscuit", "nutri choice"],
    },
    {
        "owner": "Hindustan Unilever",
        "sector": "FMCG",
        "brands": ["hindustan unilever", "surf excel", "dove", "lifebuoy", "bru"],
    },
    {
        "owner": "Nestlé India",
        "sector": "FMCG",
        "brands": ["nestle india", "nestlé india", "maggi", "nescafe", "nescafé", "kitkat", "kit kat", "munch chocolate"],
    },
    {
        "owner": "Procter & Gamble",
        "sector": "FMCG",
        "brands": ["procter & gamble", "procter and gamble", "head & shoulders", "head and shoulders", "pampers", "gillette", "oral-b"],
    },
    {
        "owner": "PepsiCo",
        "sector": "FMCG",
        "brands": ["pepsico", "pepsi co", "lay's", "lays chips", "kurkure", "quaker oats"],
    },
    {
        "owner": "Coca-Cola",
        "sector": "FMCG",
        "brands": ["coca-cola", "coca cola", "cocacola", "thums up", "thumbs up cola", "maaza", "minute maid"],
    },
    {
        "owner": "Mondelez India",
        "sector": "FMCG",
        "brands": ["mondelez india", "mondelēz india", "cadbury dairy milk", "cadbury dairymilk", "bournvita", "oreo biscuit"],
    },
    {
        "owner": "Dabur India",
        "sector": "FMCG",
        "brands": ["dabur india", "dabur", "dabur red paste", "hajmola", "vatika", "odonil"],
    },
    {
        "owner": "Marico",
        "sector": "FMCG",
        "brands": ["marico", "parachute coconut oil", "saffola", "livon serum", "set wet"],
    },
    {
        "owner": "Godrej Consumer Products",
        "sector": "FMCG",
        "brands": ["godrej consumer", "godrej aer", "goodknight", "good knight", "cinthol", "expert rich creme"],
    },
    {
        "owner": "Reckitt",
        "sector": "FMCG",
        "brands": ["reckitt", "reckitt benckiser", "dettol", "harpic", "lizol", "mortein"],
    },
    {
        "owner": "Colgate-Palmolive",
        "sector": "FMCG",
        "brands": ["colgate-palmolive", "colgate palmolive", "colgate toothpaste", "palmolive", "visible white toothpaste"],
    },
    {
        "owner": "Gujarat Cooperative Milk Marketing Federation",
        "sector": "FMCG",
        "brands": ["amul", "amul butter", "amul milk", "amulya dairy whitener"],
    },
]

EXCLUDED_ACRONYMS = {
    "PDF", "PPTX", "RFP", "RFQ", "MVP", "URL", "HTML", "HTTP", "FMCG",
    "THE", "AND", "FOR", "FROM", "WITH", "THIS", "THAT",
}

VOCAB_STOPWORDS = {
    "about", "after", "again", "against", "also", "because", "being", "between",
    "could", "their", "there", "these", "those", "through", "under", "would",
    "which", "while", "where", "what", "when", "will", "within", "without",
    "should", "shall", "must", "need", "needs", "required", "provide", "provided",
    "proposal", "project", "work", "scope", "client", "company", "organization",
    "organisation", "business", "including", "include", "services", "service",
    "deliver", "delivery", "develop", "development", "based", "make", "made",
    "into", "from", "than", "more", "very", "each", "other", "such", "have",
    "with", "that", "this", "were", "been", "they",
}


def sentences(text):
    compact = re.sub(r"\s+", " ", text).strip()
    return [item.strip() for item in re.split(r"(?<=[.!?])\s+", compact) if len(item.strip()) > 25]


def document_units(text):
    line_units = []
    for line in text.splitlines():
        clean = re.sub(r"\s+", " ", line).strip(" \t•*-")
        if len(clean) > 25:
            line_units.append(clean)
    return list(dict.fromkeys(line_units + sentences(text)))


def clause_units(units):
    clauses = []
    for item in units:
        for clause in re.split(r"(?<=[.!?])\s+|\n+|(?<=;)\s+", item):
            clean = re.sub(r"\s+", " ", clause).strip(" \t•*-")
            if len(clean) > 25 and clean not in clauses:
                clauses.append(clean)
    return clauses


def extract_objective_sections(text):
    objective_heading = (
        r"(?:case\s+|business\s+|project\s+)?objectives?|goals?|purpose|"
        r"desired\s+outcomes?|the\s+ask|what\s+we\s+need\s+to\s+achieve|"
        r"what\s+success\s+looks\s+like|success\s+criteria|expected\s+outcomes?|"
        r"(?:primary|business|key)\s+questions?|strategic\s+intent|ambition"
    )
    section_break = re.compile(
        r"^(?:\d+(?:\.\d+)*[\s.)-]*)?(?:background|context|current\s+situation|"
        r"challenges?|pain\s+points?|scope(?:\s+of\s+work)?|in\s+scope|out\s+of\s+scope|"
        r"deliverables?|requirements?|timeline|stakeholders?|approach|methodology)"
        r"\s*[:\-]?\s*$",
        flags=re.IGNORECASE,
    )
    inline_heading = re.compile(
        rf"^(?:\d+(?:\.\d+)*[\s.)-]*)?(?:{objective_heading})\s*[:\-]\s*(?P<value>.+)$",
        flags=re.IGNORECASE,
    )
    standalone_heading = re.compile(
        rf"^(?:\d+(?:\.\d+)*[\s.)-]*)?(?:{objective_heading})\s*[:\-]?\s*$",
        flags=re.IGNORECASE,
    )

    lines = [re.sub(r"\s+", " ", line).strip(" \t#*•") for line in text.splitlines()]
    objectives = []
    for index, line in enumerate(lines):
        if not line:
            continue
        inline_match = inline_heading.match(line)
        if inline_match:
            value = inline_match.group("value").strip(" \t•*-")
            if len(value) > 5 and value not in objectives:
                objectives.append(value[:420])
            continue
        if not standalone_heading.match(line):
            continue

        for following in lines[index + 1:index + 8]:
            clean = following.strip(" \t•*-")
            if not clean:
                continue
            if section_break.match(clean.split(":")[0]) or standalone_heading.match(clean):
                break
            if len(clean) > 5 and clean not in objectives:
                objectives.append(clean[:420])
            if len(objectives) >= 5:
                break

    return objectives


def first_matching(items, terms, fallback=""):
    for item in items:
        lowered = item.lower()
        if any(term in lowered for term in terms):
            return item[:420]
    return fallback


def matching_items(items, terms, limit=5):
    matches = []
    for item in items:
        lowered = item.lower()
        if any(term in lowered for term in terms) and item not in matches:
            matches.append(item[:320])
        if len(matches) >= limit:
            break
    return matches

def extract_pdf_visual_text(file_bytes, max_pages=8):
    if os.environ.get("ENABLE_PDF_OCR", "").lower() not in {"1", "true", "yes"}:
        return ""

    import pymupdf
    from rapidocr_onnxruntime import RapidOCR

    ocr_engine = RapidOCR()
    document = pymupdf.open(stream=file_bytes, filetype="pdf")
    visual_chunks = []
    try:
        for page in list(document)[:max_pages]:
            if not page.get_images(full=True):
                continue
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
            result, _ = ocr_engine(pixmap.tobytes("png"))
            if result:
                visual_chunks.extend(
                    item[1].strip()
                    for item in result
                    if len(item) > 1 and isinstance(item[1], str) and item[1].strip()
                )
    finally:
        document.close()
    return "\n".join(visual_chunks)
def extract_file_text(uploaded_file):
    text, _ = extract_file_content(uploaded_file)
    return text


def find_brand_signals(text, filename=""):
    combined = f"{filename}\n{text}".lower()
    signals = []
    for owner_record in KNOWN_BRAND_OWNERS:
        matches = [
            brand
            for brand in owner_record["brands"]
            if re.search(rf"(?<![a-z0-9]){re.escape(brand)}(?![a-z0-9])", combined)
        ]
        if matches:
            signals.append(
                {
                    "owner": owner_record["owner"],
                    "sector": owner_record["sector"],
                    "brands": matches,
                }
            )
    return signals


def infer_sector(text, brand_signals=None):
    lowered = text.lower()
    sector_terms = {
        "FMCG": [
            "fmcg", "fast-moving consumer", "consumer goods", "packaged food",
            "food product", "atta", "biscuit", "beverage", "snack", "retail",
            "sku", "distribution", "distributor", "outlet", "modern trade",
            "general trade", "consumer brand",
        ],
        "Higher education": ["university", "college", "student", "faculty", "campus"],
        "Procurement and operations": ["procurement", "spend", "sourcing", "supplier", "category playbook"],
        "Government and public sector": ["government", "state agency", "commonwealth", "department", "public sector"],
        "Healthcare": ["health", "hospital", "patient", "medicaid", "clinical"],
        "Technology and digital services": ["digital", "software", "technology", "platform", "data"],
        "Financial services": ["bank", "insurance", "financial services", "lending"],
        "Retail and e-commerce": ["e-commerce", "ecommerce", "marketplace", "retailer", "store network", "merchandising"],
        "Manufacturing": ["manufacturing", "factory", "plant", "production line", "industrial", "yield", "throughput"],
        "Automotive and mobility": ["automotive", "vehicle", "mobility", "dealer network", "fleet", "ev charging"],
        "Telecommunications": ["telecom", "telecommunications", "subscriber", "network coverage", "broadband", "5g"],
        "Energy and utilities": ["energy", "utility", "power generation", "electricity", "renewable", "grid"],
        "Transportation and logistics": ["logistics", "transportation", "freight", "warehouse", "last mile", "supply chain"],
        "Media and entertainment": ["media", "entertainment", "audience", "streaming", "content studio", "broadcast"],
        "Real estate and construction": ["real estate", "construction", "property", "developer", "infrastructure project"],
        "Travel and hospitality": ["travel", "hospitality", "hotel", "airline", "tourism", "guest experience"],
        "Professional services": ["consulting", "legal services", "accounting", "advisory", "professional services"],
        "Nonprofit and social impact": ["nonprofit", "non-profit", "ngo", "foundation", "social impact", "beneficiaries"],
        "Agriculture and agribusiness": ["agriculture", "agribusiness", "farmer", "crop", "agri-tech", "farm"],
    }
    scores = {
        sector: sum(lowered.count(term) for term in terms)
        for sector, terms in sector_terms.items()
    }
    if brand_signals:
        scores[brand_signals[0]["sector"]] += 6
    if re.search(r"\b(?:industry|sector)\s*[:\-]\s*fmcg\b", lowered):
        scores["FMCG"] += 10
    best_sector, best_score = max(scores.items(), key=lambda item: item[1])
    return best_sector if best_score else "General professional services"


def infer_client(text, filename="", brand_signals=None):
    if brand_signals:
        return brand_signals[0]["owner"]

    source = f"{filename}\n{text}"
    patterns = [
        r"(?:client|issued by|issuing organization|issuing organisation|agency|organization|organisation|prepared for)\s*[:\-]\s*([^\n|,]{3,100})",
        r"\bfor\s+([A-Z][A-Za-z&., ]{2,80})(?:\s+(?:to|on|and)\b|[.!?]|$)",
        r"(University of [A-Z][A-Za-z ,&-]+)",
        r"(State of [A-Z][A-Za-z ]+)",
        r"(Commonwealth of [A-Z][A-Za-z ]+)",
        r"([A-Z][A-Za-z& ]+ Department of [A-Z][A-Za-z& ]+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, source, flags=re.IGNORECASE)
        if match:
            return re.sub(r"\s+", " ", match.group(1)).strip(" .-")[:120]

    filename_stem = re.sub(r"[_\-]+", " ", os.path.splitext(os.path.basename(filename))[0])
    for acronym in re.findall(r"(?<![A-Za-z])[A-Z][A-Z0-9&]{1,8}(?![A-Za-z])", filename_stem):
        if acronym not in EXCLUDED_ACRONYMS:
            return acronym

    for line in text.splitlines():
        clean = re.sub(r"\s+", " ", line).strip()
        if 5 < len(clean) < 100 and not clean.lower().startswith(
            ("request for", "rfp", "proposal", "case study", "brief")
        ):
            return clean
    return ""


def extract_geography(text):
    lowered = text.lower()
    found = []
    evidence = []
    geographic_terms = [
        (r"\bindia\b", "India"),
        (r"\b(?:metro|metropolitan)\s+cities?\b", "Metro cities"),
        (r"\btier[-\s]?1\s+cities?\b", "Tier 1 cities"),
        (r"\btier[-\s]?2\s+cities?\b", "Tier 2 cities"),
        (r"\bpan[-\s]?india\b", "Pan-India"),
        (r"\bnational(?:ly)?\b", "National"),
        (r"\bglobal\b", "Global"),
    ]
    for pattern, label in geographic_terms:
        match = re.search(pattern, lowered)
        if match:
            found.append(label)
            evidence.append(match.group(0))

    city_names = [
        "Mumbai", "Delhi", "Bengaluru", "Bangalore", "Chennai", "Hyderabad",
        "Kolkata", "Pune", "Ahmedabad", "Gurugram", "Gurgaon", "Noida",
    ]
    for city in city_names:
        if re.search(rf"\b{re.escape(city)}\b", text, flags=re.IGNORECASE):
            found.append(city)
            evidence.append(city)
    for state in STATE_NAMES:
        if re.search(rf"\b{re.escape(state)}\b", text, flags=re.IGNORECASE):
            found.append(state)
            evidence.append(state)

    return "; ".join(dict.fromkeys(found)), list(dict.fromkeys(evidence))


def extract_forcing_event(items):
    strong_trigger_terms = [
        "why now", "urgent", "urgency", "deadline", "mandate", "market shift",
        "regulatory change", "new regulation", "declining", "decline", "rising",
        "accelerating", "erosion", "eroding", "shortfall", "launch deadline",
        "act now", "before the next", "must respond", "time-sensitive",
        "immediate", "upcoming", "recent change", "new entrant",
    ]
    excluded_terms = [
        "deliverable", "scope includes", "scope of work", "proposal should",
        "team will", "we will", "shall provide", "submit", "submission",
    ]
    candidates = []
    for item in items:
        lowered = item.lower()
        if any(term in lowered for term in excluded_terms):
            continue
        score = sum(3 for term in strong_trigger_terms if term in lowered)
        if "why now" in lowered or "urgency" in lowered or "deadline" in lowered:
            score += 6
        if re.search(r"\b(?:by|before|within)\s+(?:q[1-4]|\d{4}|the next|\d+\s+(?:days?|weeks?|months?))\b", lowered):
            score += 5
        if any(term in lowered for term in ["declining", "rising", "accelerating", "erosion", "shortfall"]):
            score += 2
        if score:
            candidates.append((score, item[:420]))
    candidates.sort(key=lambda item: item[0], reverse=True)
    return candidates[0][1] if candidates else ""


def extract_pain_points(items):
    pain_terms = [
        "lack of", "unable to", "not able", "inconsistent", "fragmented",
        "inefficient", "declining", "low ", "limited ", "gap", "challenge",
        "barrier", "problem", "pressure", "unmet", "complex", "manual",
        "slow ", "unpredictable", "underpenetrated", "no clear", "unclear",
        "disconnected", "volatile", "private label", "private labels", "private-label",
        "own-label", "own label", "store brand", "me-too", "me too",
        "commodit", "price pressure", "competitive intensity", "market share",
        "eroding", "erosion", "threat",
    ]
    excluded_terms = [
        "deliverable", "scope includes", "proposal", "team will", "we will",
        "shall provide", "must build", "must deliver", "should include",
    ]
    candidates = []
    for item in items:
        lowered = item.lower()
        if any(term in lowered for term in excluded_terms):
            continue
        score = sum(2 for term in pain_terms if term in lowered)
        if score:
            candidates.append((score, item[:420]))
    candidates.sort(key=lambda item: item[0], reverse=True)
    selected = []
    for _, item in candidates:
        normalized = re.sub(r"\s+", " ", item).strip().lower()
        overlapping_index = next(
            (
                index
                for index, existing in enumerate(selected)
                if normalized in existing.lower() or existing.lower() in normalized
            ),
            None,
        )
        if overlapping_index is not None:
            if len(item) < len(selected[overlapping_index]):
                selected[overlapping_index] = item
            continue
        selected.append(item)
        if len(selected) >= 5:
            break
    return selected


def extract_scope_in(items, text):
    scope_items = []
    scope_sentence_pattern = re.compile(
        r"\b(?:scope\s+(?:includes?|covers?|consists\s+of|encompasses)|"
        r"(?:work|project|engagement)\s+includes?|"
        r"responsible\s+for|services?\s+include|will\s+include|"
        r"in\s+scope(?:\s+(?:are|is))?)\s*[:\-]?\s*(?P<value>.+)",
        flags=re.IGNORECASE,
    )
    labeled_scope_pattern = re.compile(
        r"^(?P<label>products?|categories?|channels?|markets?|geograph(?:y|ies)|"
        r"customer\s+segments?|capabilities?|workstreams?|functions?|departments?|"
        r"processes?|systems?|services?)\s*[:\-]\s*(?P<value>.+)$",
        flags=re.IGNORECASE,
    )

    for item in items:
        clean = re.sub(r"\s+", " ", item).strip()
        labeled_match = labeled_scope_pattern.search(clean)
        if labeled_match:
            label = labeled_match.group("label").strip().capitalize()
            value = labeled_match.group("value").strip(" .;:-")
            if value:
                scope_items.append(f"{label}: {value}"[:320])
            continue

        sentence_match = scope_sentence_pattern.search(clean)
        if sentence_match:
            value = sentence_match.group("value").strip(" .;:-")
            if value:
                scope_items.append(value[:320])

    return list(dict.fromkeys(scope_items))[:8]


def extract_vocabulary(text):
    lowered = text.lower()
    industry_phrases = [
        "modern trade", "general trade", "packaged foods", "metro cities",
        "consumer segmentation", "brand positioning", "category growth",
        "point of purchase", "digital-first brands", "market share",
        "distribution visibility", "consumer awareness",
        "operating model", "target operating model", "service delivery",
        "customer experience", "user experience", "digital transformation",
        "artificial intelligence", "machine learning", "data platform",
        "clinical pathway", "patient outcomes", "care delivery",
        "credit risk", "claims processing", "assets under management",
        "supply chain", "last mile", "inventory turnover", "plant utilization",
        "network coverage", "subscriber growth", "renewable energy",
        "student outcomes", "faculty workload", "public service",
        "cost savings", "revenue growth", "market penetration",
    ]
    recognized_phrases = [phrase for phrase in industry_phrases if phrase in lowered]
    acronym_candidates = [
        acronym
        for acronym in re.findall(r"\b[A-Z][A-Z0-9&/-]{2,}\b", text)
        if acronym not in EXCLUDED_ACRONYMS
    ]
    words = [
        word.lower()
        for word in re.findall(r"[A-Za-z][A-Za-z0-9&/-]{3,}", text)
        if word.lower() not in VOCAB_STOPWORDS
    ]
    repeated_words = [word for word, count in Counter(words).most_common() if count >= 2]
    normalized_words = re.findall(r"[A-Za-z][A-Za-z0-9&/-]*", lowered)
    bigrams = [
        f"{left} {right}"
        for left, right in zip(normalized_words, normalized_words[1:])
        if len(left) >= 4
        and len(right) >= 4
        and left not in VOCAB_STOPWORDS
        and right not in VOCAB_STOPWORDS
    ]
    repeated_phrases = [
        phrase for phrase, count in Counter(bigrams).most_common()
        if count >= 2
    ]
    return list(
        dict.fromkeys(recognized_phrases + acronym_candidates + repeated_phrases[:4] + repeated_words[:8])
    )[:12]


def extract_quantified_targets(items):
    number_pattern = re.compile(
        r"(?:[$₹]\s?[\d,.]+\s*(?:million|billion|crore|lakh|m|b)?|"
        r"\b(?:INR|USD|EUR|GBP)\s?[\d,.]+\s*(?:million|billion|crore|lakh|m|b)?|"
        r"\b\d+(?:\.\d+)?%(?!\w)|\b\d+(?:\.\d+)?x\b|"
        r"\b\d+\s*(?:million|billion|crore|lakh)\b)",
        flags=re.IGNORECASE,
    )
    target_terms = [
        "aim", "target", "objective", "goal", "commit", "expected to",
        "seeks to", "plans to", "increase", "reduce", "decrease", "improve",
        "achieve", "deliver", "save", "reach", "raise", "lower", "cut",
        "grow", "at least", "no more than", "maximum", "minimum",
    ]
    historical_terms = [
        "currently", "today", "last year", "historically", "accounts for",
        "represents", "survey", "respondents", "existing baseline",
        "market currently", "as of", "target audience", "target consumers",
        "target segment",
    ]
    candidates = []
    for item in items:
        if not number_pattern.search(item):
            continue
        lowered = item.lower()
        if any(term in lowered for term in ["target audience", "target consumers", "target segment"]) and not any(
            action in lowered
            for action in ["increase", "reduce", "improve", "achieve", "reach", "raise", "lower", "grow"]
        ):
            continue
        score = sum(3 for term in target_terms if term in lowered)
        if any(term in lowered for term in ["target", "objective", "goal", "at least", "no more than"]):
            score += 4
        if any(term in lowered for term in historical_terms):
            score -= 4
        if score >= 3:
            candidates.append((score, item[:420]))
    candidates.sort(key=lambda candidate: candidate[0], reverse=True)
    return list(dict.fromkeys(item for _, item in candidates[:3]))


def extract_case_objectives(items):
    objective_terms = [
        "objective is", "objectives are", "goal is", "goals are", "aims to",
        "aim is to", "seeks to", "wants to", "intends to", "purpose is",
        "designed to", "needs to", "need to", "desired outcome",
        "business outcome", "case objective", "should enable", "objective:",
        "the objective", "the aim", "case is to",
        "project is to", "success means", "looking to", "we need to",
        "aim to", "seek to", "want to", "intend to", "ambition is",
    ]
    excluded_terms = [
        "deliverable", "scope includes", "scope of work", "proposal should",
        "team will", "we will", "shall provide", "submit", "submission",
        "vendor must", "consultant must", "scope covers", "in scope",
        "out of scope", "with a focus on", "geographic focus", "channel focus",
    ]
    candidates = []
    for item in items:
        lowered = item.lower()
        if any(term in lowered for term in excluded_terms):
            continue
        if any(term in lowered for term in ["need to act now", "needs to act now", "must act now"]):
            continue
        score = sum(3 for term in objective_terms if term in lowered)
        if re.match(
            r"^(?:improve|reduce|increase|enable|establish|achieve|accelerate|"
            r"strengthen|streamline|eliminate|optimize|optimise)\b",
            lowered,
        ):
            score += 3
        if any(term in lowered for term in ["objective is", "goal is", "purpose is", "desired outcome"]):
            score += 4
        if score:
            candidates.append((score, item[:420]))
    candidates.sort(key=lambda candidate: candidate[0], reverse=True)
    return list(dict.fromkeys(item for _, item in candidates[:3]))


def digest_text(text, filename, visual_text=""):
    text = text[:250000]
    visual_text = visual_text[:50000]
    analysis_text = "\n".join(part for part in [text, visual_text] if part.strip())
    items = clause_units(document_units(analysis_text))
    text_brand_signals = find_brand_signals(text)
    visual_brand_signals = find_brand_signals(visual_text) if not text_brand_signals else []
    filename_brand_signals = (
        find_brand_signals("", filename)
        if not text_brand_signals and not visual_brand_signals
        else []
    )
    brand_signals = text_brand_signals or visual_brand_signals or filename_brand_signals
    geography, geography_evidence = extract_geography(analysis_text)
    quantified_targets = extract_quantified_targets(items)
    section_objectives = extract_objective_sections(analysis_text)
    inferred_objectives = extract_case_objectives(items)
    objective_statements = list(dict.fromkeys(section_objectives + inferred_objectives))[:5]
    target_type = (
        "mixed"
        if quantified_targets and objective_statements
        else "quantified"
        if quantified_targets
        else "qualitative objective"
        if objective_statements
        else "not identified"
    )
    timeline_matches = re.findall(
        r"\b(?:\d{1,3}\s*(?:business\s+)?(?:days?|weeks?|months?|years?)|"
        r"(?:January|February|March|April|May|June|July|August|September|October|November|December)"
        r"\s+\d{1,2}(?:,\s+\d{4})?)\b",
        analysis_text,
        flags=re.IGNORECASE,
    )
    stakeholder_terms = [
        "executive leadership", "agency leaders", "stakeholders", "faculty",
        "students", "employees", "vendors", "suppliers", "project manager",
        "steering committee", "community", "residents", "customers",
        "consumers", "distributors", "retailers",
    ]
    lowered = analysis_text.lower()
    stakeholders = [term.title() for term in stakeholder_terms if term in lowered]
    scope_in = extract_scope_in(items, analysis_text)
    scope_out = matching_items(
        items,
        ["out of scope", "excluded", "will not include", "not include", "outside the scope"],
        6,
    )
    deliverables = matching_items(
        items,
        ["deliverable", "roadmap", "model", "playbook", "database", "report", "implementation plan", "recommendation"],
        6,
    )
    client = infer_client(analysis_text, filename, brand_signals)
    sector = infer_sector(analysis_text, brand_signals)
    forcing_event = extract_forcing_event(items)
    pain_points = extract_pain_points(items)
    objective_status = "extracted"
    if not objective_statements:
        if quantified_targets:
            objective_statements = quantified_targets[:3]
        elif pain_points:
            objective_statements = ["Resolve the stated problem: " + pain_points[0]]
            objective_status = "proposed"
        else:
            objective_status = "needs input"
    if objective_statements:
        target_type = "mixed" if quantified_targets else "qualitative objective"

    return {
        "client": client,
        "sector": sector,
        "geography": geography,
        "forcing_event": forcing_event,
        "case_objective": " ".join(objective_statements[:3]),
        "objective_status": objective_status,
        "objective_requires_confirmation": True,
        "quantified_target": " ".join(quantified_targets[:3]),
        "target_type": target_type,
        "scope_in": scope_in,
        "scope_out": scope_out,
        "timeline": "; ".join(dict.fromkeys(timeline_matches[:4])),
        "stakeholders": stakeholders[:8],
        "deliverables": deliverables,
        "jargon_terms": extract_vocabulary(analysis_text),
        "pain_points": pain_points,
        "pain_points_inferred": True,
        "brand_signals": brand_signals,
        "client_evidence": [
            (
                f"{signal['owner']} inferred from visual evidence in the PDF "
                f"(logo text: {', '.join(signal['brands'])})"
                if visual_brand_signals
                else (
                    f"{signal['owner']} inferred from the source filename "
                    f"(brand text: {', '.join(signal['brands'])})"
                    if filename_brand_signals
                    else f"{signal['owner']} inferred from sub-brand(s): {', '.join(signal['brands'])}"
                )
            )
            for signal in brand_signals
        ],
        "client_inferred_from_visual_evidence": bool(visual_brand_signals),
        "sector_evidence": [
            term for term in ["FMCG", "consumer goods", "packaged food", "atta", "retail", "SKU"]
            if term.lower() in lowered
        ],
        "geography_evidence": geography_evidence,
        "forcing_event_evidence": forcing_event,
        "case_objective_evidence": pain_points[:1] if objective_status == "proposed" else objective_statements,
        "quantified_target_evidence": quantified_targets,
        "pain_point_evidence": pain_points,
        "source_file": filename,
        "source_character_count": len(analysis_text),
    }


def normalized_evidence_supported(evidence, source_text):
    if not evidence:
        return False
    normalized_source = re.sub(r"\s+", " ", source_text).lower()
    snippets = evidence if isinstance(evidence, list) else [evidence]
    for snippet in snippets:
        if not isinstance(snippet, str):
            continue
        normalized_snippet = re.sub(r"\s+", " ", snippet).strip().lower()
        if len(normalized_snippet) >= 6 and normalized_snippet in normalized_source:
            return True
        snippet_words = set(re.findall(r"[a-z0-9]+", normalized_snippet))
        if len(snippet_words) >= 3:
            source_words = set(re.findall(r"[a-z0-9]+", normalized_source))
            if len(snippet_words & source_words) / len(snippet_words) >= 0.8:
                return True
    return False


def value_supported_by_evidence(value, evidence, source_text):
    if not isinstance(value, str) or not value.strip():
        return False
    snippets = evidence if isinstance(evidence, list) else [evidence]
    value_normalized = re.sub(r"\s+", " ", value).strip().lower()
    value_numbers = set(re.findall(r"(?:[$₹]\s*)?\d+(?:\.\d+)?%?", value_normalized))
    ignored_words = {
        "a", "an", "and", "by", "for", "from", "in", "into", "of", "on",
        "or", "the", "to", "within", "with",
    }
    value_words = {
        word for word in re.findall(r"[a-z0-9]+", value_normalized)
        if word not in ignored_words and len(word) > 1
    }
    for snippet in snippets:
        if not isinstance(snippet, str) or not normalized_evidence_supported(snippet, source_text):
            continue
        snippet_normalized = re.sub(r"\s+", " ", snippet).strip().lower()
        if value_normalized in snippet_normalized:
            return True
        snippet_numbers = set(re.findall(r"(?:[$₹]\s*)?\d+(?:\.\d+)?%?", snippet_normalized))
        if value_numbers and not value_numbers.issubset(snippet_numbers):
            continue
        snippet_words = {
            word for word in re.findall(r"[a-z0-9]+", snippet_normalized)
            if word not in ignored_words and len(word) > 1
        }
        if value_words and value_words.issubset(snippet_words):
            return True
    return False


def forcing_event_has_trigger(value, evidence=None):
    combined = " ".join(
        [str(value or "")]
        + ([str(item) for item in evidence] if isinstance(evidence, list) else [str(evidence or "")])
    ).lower()
    return bool(
        any(
            term in combined
            for term in [
                "why now", "urgent", "urgency", "deadline", "mandate",
                "market shift", "regulatory change", "new regulation",
                "declining", "decline", "rising", "accelerating", "erosion",
                "eroding", "shortfall", "time-sensitive", "immediate",
                "upcoming", "new entrant", "must respond", "act now",
            ]
        )
        or re.search(
            r"\b(?:by|before|within)\s+(?:q[1-4]|\d{4}|the next|\d+\s+(?:days?|weeks?|months?))\b",
            combined,
        )
    )


def interpret_digest_with_groq(text, visual_text, filename, baseline):
    api_key = os.environ.get("GROQ_API_KEY")
    if not api_key:
        app.logger.info("Groq interpretation skipped: GROQ_API_KEY is not configured")
        baseline["interpretation_source"] = "deterministic fallback"
        baseline["interpretation_warning"] = "Groq is not configured; deterministic extraction was used."
        return baseline

    source_text = "\n".join(part for part in [text, visual_text] if part.strip())
    if len(source_text) > 80000:
        source_text = source_text[:60000] + "\n[...middle omitted...]\n" + source_text[-20000:]

    baseline_fields = {
        key: baseline.get(key)
        for key in [
            "client", "sector", "geography", "forcing_event", "case_objective",
            "quantified_target", "target_type", "scope_in", "scope_out",
            "timeline", "stakeholders", "deliverables", "jargon_terms", "pain_points",
        ]
    }
    schema_example = {
        "client": "",
        "sector": "",
        "geography": "",
        "forcing_event": "",
        "case_objective": "",
        "quantified_target": "",
        "target_type": "quantified|qualitative objective|mixed|not identified",
        "scope_in": [],
        "scope_out": [],
        "timeline": "",
        "stakeholders": [],
        "deliverables": [],
        "jargon_terms": [],
        "pain_points": [],
        "field_evidence": {
            "client": [],
            "sector": [],
            "geography": [],
            "forcing_event": [],
            "case_objective": [],
            "quantified_target": [],
            "scope_in": [],
            "scope_out": [],
            "timeline": [],
            "stakeholders": [],
            "deliverables": [],
            "jargon_terms": [],
            "pain_points": [],
        },
    }
    prompt = f"""
Interpret the case/RFP text into the exact JSON structure below.

Field meanings:
- client: the parent organisation that owns the problem, not a vendor or product. Infer a parent from a clearly cited sub-brand only when justified.
- sector: the client's industry, not the project workstream.
- geography: all material countries, regions, city tiers, and city focus areas.
- forcing_event: only an explicit trigger, deadline, urgency, material recent change, or consequence that explains why action is needed now. Leave blank rather than restating a pain point or generic requirement.
- case_objective: the intended qualitative change or business outcome. Populate this even when there is also a quantified target.
- quantified_target: only an explicit measurable target stated as an intended outcome. Leave blank when the source has no such target.
- target_type: one of quantified, qualitative objective, mixed, or not identified.
- scope_in: products, services, categories, channels, customer groups, capabilities, systems, processes, functions, departments, workstreams, and geographies explicitly in scope.
- scope_out: only explicit exclusions.
- pain_points: negative current conditions, barriers, threats, or friction. Include competitive threats such as private labels when stated.
- jargon_terms: repeated case language and industry-specific terminology, not generic document words.

Rules:
1. Do not invent facts or numbers.
2. A descriptive statistic is not automatically a target.
3. Return concise field values, not commentary about your reasoning.
4. For every populated field, include one or more short source excerpts in field_evidence.
5. Evidence must be copied from the supplied source closely enough to verify.
6. Return JSON only, matching this structure:
{json.dumps(schema_example)}

Filename:
{filename}

Deterministic first-pass candidates (correct them when the source supports a better interpretation):
{json.dumps(baseline_fields, ensure_ascii=False)}

Extracted source text:
{source_text}
"""
    try:
        from groq import Groq

        model = os.environ.get("GROQ_MODEL", "openai/gpt-oss-120b")
        app.logger.info(
            "Groq interpretation request started: model=%s source_chars=%d",
            model,
            len(source_text),
        )
        client = Groq(api_key=api_key, timeout=25, max_retries=1)
        completion = client.chat.completions.create(
            model=model,
            messages=[
                {
                    "role": "system",
                    "content": "You are a precise business-case analyst. Return valid JSON only and ground every field in source evidence.",
                },
                {"role": "user", "content": prompt},
            ],
            temperature=0,
            response_format={"type": "json_object"},
        )
        interpreted = json.loads(completion.choices[0].message.content)
        field_evidence = interpreted.get("field_evidence") or {}
        merged = deepcopy(baseline)

        string_fields = [
            "client", "sector", "geography", "forcing_event", "case_objective",
            "quantified_target", "timeline",
        ]
        list_fields = [
            "scope_in", "scope_out", "stakeholders", "deliverables",
            "jargon_terms", "pain_points",
        ]
        for field in string_fields:
            value = interpreted.get(field)
            if (
                field == "forcing_event"
                and isinstance(value, str)
                and value.strip()
                and not forcing_event_has_trigger(value, field_evidence.get(field))
            ):
                continue
            if (
                field == "quantified_target"
                and isinstance(value, str)
                and value.strip()
                and not re.search(
                    r"(?:[$₹]\s?[\d,.]+|\b(?:INR|USD|EUR|GBP)\b|"
                    r"\b\d+(?:\.\d+)?%|\b\d+(?:\.\d+)?x\b|"
                    r"\b\d+\s*(?:million|billion|crore|lakh)\b)",
                    value,
                    flags=re.IGNORECASE,
                )
            ):
                continue
            if (
                isinstance(value, str)
                and value.strip()
                and value_supported_by_evidence(
                    value, field_evidence.get(field), source_text
                )
            ):
                merged[field] = value.strip()
                if field == "case_objective":
                    merged["objective_status"] = "extracted"
        for field in list_fields:
            value = interpreted.get(field)
            evidence = field_evidence.get(field)
            if isinstance(value, list):
                cleaned = [
                    str(item).strip()
                    for item in value
                    if str(item).strip()
                    and value_supported_by_evidence(str(item), evidence, source_text)
                ]
                if cleaned:
                    if field in {"scope_in", "scope_out", "pain_points"}:
                        combined = list(baseline.get(field) or [])
                        for item in cleaned:
                            normalized = re.sub(r"\s+", " ", item).strip().lower()
                            if not any(
                                normalized in re.sub(r"\s+", " ", existing).strip().lower()
                                or re.sub(r"\s+", " ", existing).strip().lower() in normalized
                                for existing in combined
                            ):
                                combined.append(item)
                        merged[field] = combined[:10]
                    else:
                        merged[field] = cleaned[:10]

        has_quantified_target = bool(merged.get("quantified_target"))
        has_case_objective = bool(merged.get("case_objective"))
        merged["target_type"] = (
            "mixed"
            if has_quantified_target and has_case_objective
            else "quantified"
            if has_quantified_target
            else "qualitative objective"
            if has_case_objective
            else "not identified"
        )
        merged["groq_field_evidence"] = field_evidence
        merged["interpretation_source"] = "Groq structured interpretation with deterministic fallback"
        merged["interpretation_model"] = model
        merged["interpretation_request_id"] = getattr(completion, "id", "") or ""
        merged["interpretation_warning"] = ""
        app.logger.info(
            "Groq interpretation request completed: model=%s request_id=%s",
            model,
            merged["interpretation_request_id"] or "unavailable",
        )
        return merged
    except Exception as error:
        app.logger.warning("Groq interpretation failed; deterministic fallback used (%s)", type(error).__name__)
        baseline["interpretation_source"] = "deterministic fallback"
        baseline["interpretation_warning"] = "Groq interpretation was unavailable; deterministic extraction was used."
        return baseline


def select_shape(digest):
    signal_text = " ".join(
        [
            str(digest.get("sector", "")),
            str(digest.get("forcing_event", "")),
            str(digest.get("case_objective", "")),
            str(digest.get("quantified_target", "")),
            str(digest.get("timeline", "")),
            " ".join(digest.get("scope_in") or []),
            " ".join(digest.get("scope_out") or []),
            " ".join(digest.get("deliverables") or []),
            " ".join(digest.get("pain_points") or []),
            " ".join(digest.get("jargon_terms") or []),
        ]
    ).lower()

    if any(term in signal_text for term in ["procurement", "spend cube", "sourcing", "category playbook", "supplier"]):
        return STRUCTURE_LIBRARY[5]
    if any(term in signal_text for term in ["accessibility", "inclusive design", "visual designer", "user experience", "ux"]):
        return STRUCTURE_LIBRARY[0]
    if any(term in signal_text for term in ["training", "coaching", "capability building", "capacity transfer", "roster"]):
        return STRUCTURE_LIBRARY[2]
    if "education" in signal_text or any(term in signal_text for term in ["university", "college", "campus", "student"]):
        if any(term in signal_text for term in ["cost reduction", "shortfall", "savings", "operating expense"]) or digest.get("scope_out"):
            return STRUCTURE_LIBRARY[4]
        return STRUCTURE_LIBRARY[3]
    if any(term in signal_text for term in ["scenario model", "policy design", "impact model", "readiness assessment"]):
        return STRUCTURE_LIBRARY[3]
    if digest.get("scope_out") and any(term in signal_text for term in ["cost", "efficiency", "savings"]):
        return STRUCTURE_LIBRARY[4]
    return STRUCTURE_LIBRARY[1]


def build_draft(digest, shape):
    client = digest.get("client") or "the client"
    objective = digest.get("case_objective") or digest.get("quantified_target") or ""
    quantified_target = digest.get("quantified_target") or ""
    pain_points = [item for item in (digest.get("pain_points") or []) if item]
    scope_in = [item for item in (digest.get("scope_in") or []) if item]
    scope_out = [item for item in (digest.get("scope_out") or []) if item]
    forcing_event = digest.get("forcing_event") or ""
    problem_statement = (
        " ".join(pain_points[:3])
        or forcing_event
        or f"{client} is seeking a solution aligned to the confirmed objective and scope."
    )
    scope_summary = "; ".join(scope_in[:4]) or "the confirmed in-scope areas"
    headline = (
        f"A focused path to {objective.rstrip('.')}"
        if objective
        else f"A problem-led solution for {client}"
    )
    phases = [
        {
            "name": "Align and diagnose",
            "duration": "Weeks 1-2",
            "goal": f"Validate the problem, baseline, stakeholders, and decision criteria across {scope_summary}.",
        },
        {
            "name": "Design the response",
            "duration": "Weeks 2-4",
            "goal": f"Develop solution options explicitly tied to {objective or 'the confirmed case objective'}.",
        },
        {
            "name": "Validate and prioritize",
            "duration": "Weeks 4-6",
            "goal": "Test options against impact, feasibility, stakeholder acceptance, and implementation risk.",
        },
        {
            "name": "Mobilize and transfer",
            "duration": "Weeks 6-8",
            "goal": "Translate the preferred solution into an owned roadmap, measures, governance, and capability transfer.",
        },
    ]
    requested_deliverables = [item for item in (digest.get("deliverables") or []) if item]
    handed_over_artifact = (
        "; ".join(requested_deliverables[:3])
        if requested_deliverables
        else "Prioritized solution roadmap, implementation plan, and measurable success framework"
    )
    risk_note = (
        f"Protect the explicit scope boundaries: {'; '.join(scope_out[:3])}."
        if scope_out
        else "Validate assumptions and implementation ownership before committing to impact or timing."
    )
    return {
        "shape_id": shape["id"],
        "reference_pattern": shape["name"],
        "problem_statement": problem_statement,
        "case_objective": objective,
        "quantified_target": quantified_target,
        "why_now": forcing_event,
        "pain_points": pain_points,
        "scope_in": scope_in,
        "scope_out": scope_out,
        "approach_phases": phases,
        "quantified_headline": headline,
        "solution_summary": (
            f"Address the validated pain points through a staged solution covering {scope_summary}, "
            f"with each phase tied to {objective or 'the confirmed objective'}."
        ),
        "handed_over_artifact": handed_over_artifact,
        "evidence": [],
        "risk_note": risk_note,
        "why_us": "Why-us claims will be grounded in the supplied angle, credentials, and researched evidence; no reference-case credentials are reused.",
        "client_angle": "",
        "angle_application": [],
        "enrichment_notes": "",
    }


def is_public_url(url):
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        return False
    try:
        addresses = socket.getaddrinfo(parsed.hostname, parsed.port or (443 if parsed.scheme == "https" else 80))
    except socket.gaierror:
        return False
    for address in addresses:
        ip = ipaddress.ip_address(address[4][0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
            return False
    return True


def fetch_public_url(url):
    current = url.strip()
    headers = {"User-Agent": "CaseBuilder/1.0 (+manual user-supplied URL fetch)"}
    for _ in range(4):
        if not is_public_url(current):
            raise ValueError("Only public HTTP or HTTPS URLs can be fetched.")
        response = requests.get(current, headers=headers, timeout=10, allow_redirects=False)
        if response.is_redirect:
            current = urljoin(current, response.headers["Location"])
            continue
        response.raise_for_status()
        if len(response.content) > 3 * 1024 * 1024:
            raise ValueError("The linked page is too large to process.")
        return response, current
    raise ValueError("The linked page redirected too many times.")


def page_profile(url):
    response, final_url = fetch_public_url(url)
    soup = BeautifulSoup(response.text, "html.parser")
    title = soup.title.get_text(" ", strip=True) if soup.title else final_url
    description_tag = soup.find("meta", attrs={"name": re.compile("^description$", re.I)})
    description = description_tag.get("content", "").strip() if description_tag else ""
    for tag in soup(["script", "style", "noscript", "svg"]):
        tag.decompose()
    main = soup.find("main") or soup.find("article") or soup.body or soup
    text = re.sub(r"\s+", " ", main.get_text(" ", strip=True))
    return {
        "url": final_url,
        "title": title[:180],
        "description": description[:500],
        "excerpt": text[:1600],
    }, soup


def search_public_web(query, limit=4):
    if not query.strip():
        return []
    response = requests.get(
        f"https://html.duckduckgo.com/html/?q={quote_plus(query)}",
        headers={"User-Agent": "Mozilla/5.0 CaseBuilder/1.0"},
        timeout=12,
    )
    response.raise_for_status()
    soup = BeautifulSoup(response.text, "html.parser")
    urls = []
    for link in soup.select("a.result__a"):
        href = (link.get("href") or "").strip()
        if not href:
            continue
        parsed = urlparse(href)
        if "duckduckgo.com" in (parsed.hostname or ""):
            redirected = parse_qs(parsed.query).get("uddg", [""])[0]
            href = unquote(redirected)
        if href.startswith("//"):
            href = f"https:{href}"
        if href and is_public_url(href) and href not in urls:
            urls.append(href)
        if len(urls) >= limit:
            break
    return urls


def search_wikipedia_sources(query, limit=2):
    response = requests.get(
        "https://en.wikipedia.org/w/api.php",
        params={
            "action": "query",
            "list": "search",
            "srsearch": query,
            "srlimit": limit,
            "format": "json",
        },
        headers={"User-Agent": "CaseBuilder/1.0"},
        timeout=12,
    )
    response.raise_for_status()
    matches = response.json().get("query", {}).get("search", [])
    sources = []
    for match in matches[:limit]:
        title = match.get("title", "")
        if not title:
            continue
        extract_response = requests.get(
            "https://en.wikipedia.org/w/api.php",
            params={
                "action": "query",
                "prop": "extracts",
                "exintro": 1,
                "explaintext": 1,
                "titles": title,
                "format": "json",
            },
            headers={"User-Agent": "CaseBuilder/1.0"},
            timeout=12,
        )
        extract_response.raise_for_status()
        pages = extract_response.json().get("query", {}).get("pages", {})
        extract = next(iter(pages.values()), {}).get("extract", "")
        sources.append(
            {
                "url": f"https://en.wikipedia.org/wiki/{quote_plus(title).replace('+', '_')}",
                "title": title,
                "description": extract[:500],
                "excerpt": extract[:1600],
                "status": "fetched",
                "source_type": "web research",
                "research_query": query,
            }
        )
    return sources


def search_news_sources(query, limit=3):
    response = requests.get(
        f"https://news.google.com/rss/search?q={quote_plus(query)}&hl=en-IN&gl=IN&ceid=IN:en",
        headers={"User-Agent": "CaseBuilder/1.0"},
        timeout=12,
    )
    response.raise_for_status()
    root = ET.fromstring(response.content)
    sources = []
    for item in root.findall(".//item")[:limit]:
        title = (item.findtext("title") or "").strip()
        link = (item.findtext("link") or "").strip()
        description = BeautifulSoup(
            item.findtext("description") or "",
            "html.parser",
        ).get_text(" ", strip=True)
        if title and link:
            sources.append(
                {
                    "url": link,
                    "title": title,
                    "description": description[:500],
                    "excerpt": description[:1600],
                    "status": "fetched",
                    "source_type": "web research",
                    "research_query": query,
                }
            )
    return sources


def detect_brand(url):
    profile, soup = page_profile(url)
    html = str(soup)
    theme_tag = soup.find("meta", attrs={"name": re.compile("^theme-color$", re.I)})
    theme = theme_tag.get("content", "") if theme_tag else ""
    colors = re.findall(r"#[0-9a-fA-F]{6}\b", html)
    colors = list(dict.fromkeys(color.upper() for color in colors))
    primary = theme.upper() if re.fullmatch(r"#[0-9a-fA-F]{6}", theme or "") else (colors[0] if colors else NEUTRAL_BRAND["primary"])
    secondary = next((color for color in colors if color != primary and color not in {"#000000", "#FFFFFF"}), NEUTRAL_BRAND["secondary"])
    accent = next((color for color in colors if color not in {primary, secondary, "#000000", "#FFFFFF"}), NEUTRAL_BRAND["accent"])
    logo_url = ""
    logo = soup.find("img", attrs={"class": re.compile("logo", re.I)}) or soup.find("img", attrs={"id": re.compile("logo", re.I)})
    if logo and logo.get("src"):
        logo_url = urljoin(profile["url"], logo["src"])
    else:
        icon = soup.find("link", rel=lambda value: value and "icon" in value)
        if icon and icon.get("href"):
            logo_url = urljoin(profile["url"], icon["href"])
    return {
        "primary": primary,
        "secondary": secondary,
        "accent": accent,
        "logo_url": logo_url,
        "logo_data_url": "",
        "tone": profile["description"] or NEUTRAL_BRAND["tone"],
        "source": profile["url"],
    }


def valid_hex(value, fallback):
    return value if isinstance(value, str) and re.fullmatch(r"#[0-9a-fA-F]{6}", value) else fallback


def brand_color(value):
    return RGBColor.from_string(value.lstrip("#"))


def add_background(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_text(slide, text, left, top, width, height, color, size, bold=False):
    color = readable_color(color, text_background(slide, left, top, width, height))
    visible, overflow, size = fit_text(str(text or ""), width, height, size, bold)
    if overflow:
        if not hasattr(slide.part, "_text_overflow"):
            slide.part._text_overflow = []
        slide.part._text_overflow.append(overflow)
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = visible
    run.font.name = "DejaVu Sans"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    style_frame(frame, size)
    return box


def add_bullets(slide, items, left, top, width, height, color, size=16, limit=6):
    cleaned = [re.sub(r"\s+", " ", str(item)).strip() for item in (items or []) if str(item).strip()]
    return add_text(
        slide,
        "\n".join(f"• {item}" for item in cleaned[:limit]),
        left,
        top,
        width,
        height,
        color,
        size,
    )


def add_logo(slide, brand_kit, left, top, width):
    data_url = brand_kit.get("logo_data_url", "")
    if not data_url.startswith("data:image/"):
        return
    try:
        encoded = data_url.split(",", 1)[1]
        image = base64.b64decode(encoded, validate=True)
        if len(image) <= 2 * 1024 * 1024:
            picture = slide.shapes.add_picture(io.BytesIO(image), left, top, width=width)
            if picture.height > Inches(.95):
                ratio = Inches(.95) / picture.height
                picture.width = int(picture.width * ratio)
                picture.height = Inches(.95)
    except (ValueError, IndexError):
        return


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/library")
def library():
    return jsonify(STRUCTURE_LIBRARY)


@app.post("/api/digest")
def digest():
    try:
        uploaded_file = request.files["rfp_file"]
        text, visual_text = extract_file_content(uploaded_file)
        if not text.strip() and not visual_text.strip():
            return jsonify({"error": "No readable text was found in the file."}), 422
        filename = uploaded_file.filename or "uploaded file"
        baseline = digest_text(text, filename, visual_text)
        interpreted = interpret_digest_with_groq(text, visual_text, filename, baseline)
        return jsonify(interpreted)
    except KeyError:
        return jsonify({"error": "Upload the file using the rfp_file field."}), 400
    except (ValueError, OSError) as error:
        return jsonify({"error": str(error)}), 400


@app.post("/api/draft")
def draft():
    payload = request.get_json() or {}
    digest_data = payload.get("digest", payload)
    objective = digest_data.get("case_objective")
    if not isinstance(objective, str) or not objective.strip():
        return jsonify({"error": "A case objective is required. Enter and confirm the intended outcome before drafting."}), 422
    digest_data["case_objective"] = objective.strip()
    shape = select_shape(digest_data)
    return jsonify(build_draft(digest_data, shape))


@app.post("/api/enrich")
def enrich():
    payload = request.get_json() or {}
    draft_data = deepcopy(payload.get("draft") or {})
    notes = (payload.get("enrichment") or "").strip()
    draft_data["enrichment_notes"] = notes
    draft_data["client_angle"] = notes
    if notes:
        angle_items = [
            re.sub(r"\s+", " ", item).strip(" \t•*-")
            for item in re.split(r"\n+|(?<=[.!?])\s+", notes)
            if len(re.sub(r"\s+", " ", item).strip(" \t•*-")) > 4
        ]
        draft_data["angle_application"] = angle_items[:6]
        draft_data["solution_summary"] = (
            f"{draft_data.get('solution_summary', '').rstrip()} "
            f"The client-specific angle is applied as a design constraint: {notes}"
        ).strip()
        phases = deepcopy(draft_data.get("approach_phases") or [])
        if phases:
            phases[0]["goal"] = (
                f"{phases[0].get('goal', '').rstrip()} Confirm how the supplied angle changes priorities and success criteria."
            )
        if len(phases) > 1:
            phases[1]["goal"] = (
                f"{phases[1].get('goal', '').rstrip()} Build the supplied angle directly into the solution options."
            )
        draft_data["approach_phases"] = phases
        draft_data["why_us"] = (
            "Our response is differentiated by applying the user's explicit point of view to the problem, "
            "rather than reproducing a reference-case answer."
        )
    return jsonify(draft_data)


@app.post("/api/brand")
def brand():
    payload = request.get_json() or {}
    url = (payload.get("url") or "").strip()
    if not url:
        return jsonify(NEUTRAL_BRAND)
    try:
        return jsonify(detect_brand(url))
    except (ValueError, requests.RequestException) as error:
        fallback = deepcopy(NEUTRAL_BRAND)
        fallback["warning"] = f"Brand detection failed: {error}. Neutral palette applied."
        return jsonify(fallback)


@app.post("/api/context")
def context():
    payload = request.get_json() or {}
    supplied_urls = [
        url.strip() for url in (payload.get("urls") or [])
        if isinstance(url, str) and url.strip()
    ]
    brand_url = (payload.get("brand_url") or "").strip()
    if brand_url:
        supplied_urls.insert(0, brand_url)

    results = []
    for url in dict.fromkeys(supplied_urls):
        try:
            profile, _ = page_profile(url)
            profile["status"] = "fetched"
            profile["source_type"] = "user supplied"
            results.append(profile)
        except (ValueError, requests.RequestException) as error:
            results.append(
                {
                    "url": url,
                    "status": "failed",
                    "source_type": "user supplied",
                    "error": str(error),
                }
            )

    return jsonify({"sources": results, "research_query": ""})


def synthesize_storyline(digest_data, draft_data, successful_sources):
    source_summary = [
        {
            "title": source.get("title", ""),
            "url": source.get("url", ""),
            "description": source.get("description", ""),
            "excerpt": source.get("excerpt", "")[:900],
            "source_type": source.get("source_type", ""),
        }
        for source in successful_sources[:8]
    ]
    pain_points = list(digest_data.get("pain_points") or draft_data.get("pain_points") or [])
    objective = digest_data.get("case_objective") or draft_data.get("case_objective") or ""
    quantified_target = digest_data.get("quantified_target") or draft_data.get("quantified_target") or ""
    why_now = digest_data.get("forcing_event") or draft_data.get("why_now") or ""
    client_angle = draft_data.get("client_angle") or draft_data.get("enrichment_notes") or ""
    fallback_market_insights = [
        {
            "insight": (source.get("description") or source.get("excerpt", ""))[:320],
            "source_title": source.get("title", ""),
            "url": source.get("url", ""),
        }
        for source in source_summary
        if source.get("description") or source.get("excerpt")
    ][:4]
    fallback = {
        "headline": draft_data.get("quantified_headline", ""),
        "client": digest_data.get("client", ""),
        "executive_summary": draft_data.get("solution_summary", ""),
        "case_objective": objective,
        "quantified_target": quantified_target,
        "why_now": why_now,
        "problem_statement": draft_data.get("problem_statement") or " ".join(pain_points[:3]),
        "pain_points": pain_points,
        "scope_in": list(digest_data.get("scope_in") or []),
        "scope_out": list(digest_data.get("scope_out") or []),
        "solution_summary": draft_data.get("solution_summary", ""),
        "approach_phases": list(draft_data.get("approach_phases") or []),
        "client_angle": client_angle,
        "angle_application": list(draft_data.get("angle_application") or []),
        "why_us": draft_data.get("why_us", ""),
        "why_us_evidence": list(draft_data.get("evidence") or []),
        "handed_over_artifact": draft_data.get("handed_over_artifact", ""),
        "risk_note": draft_data.get("risk_note", ""),
        "market_insights": fallback_market_insights,
        "market_context": source_summary,
        "storyline_sections": [
            "Executive summary",
            "Case objective and why now",
            "Key pain points",
            "Our solution",
            "Our approach",
            "Why us",
            "Market evidence and sources",
        ],
        "confirmation_status": "awaiting_confirmation",
    }

    api_key = os.environ.get("GROQ_API_KEY")
    if not api_key:
        return fallback

    schema = {
        key: value
        for key, value in fallback.items()
        if key not in {"market_context", "confirmation_status"}
    }
    prompt = f"""
Build a problem-led proposal storyline from the validated case facts, the working solution,
the user's angle, and researched web sources. Return JSON matching the supplied schema.

Rules:
1. The validated digest is the source of truth for objective, why-now, pain points, scope, and target.
2. Leave why_now blank when no explicit trigger, urgency, deadline, or material change is stated.
3. Treat the reference pattern only as method inspiration. Do not reuse credentials, client names,
   statistics, claims, or answers from reference proposals.
4. Integrate the user's angle into solution_summary, angle_application, and relevant phase goals.
5. Use web research only for market/client context. Omit weakly related sources. Every market insight
   must directly inform the validated pain, objective, scope, or solution and include the source title and URL.
6. Do not invent firm credentials. why_us may use only supplied angle/credentials and the demonstrated fit
   of the proposed approach to this case. If credentials are absent, make that limitation explicit.
7. Preserve at least four practical approach phases and keep every phase tied to the validated problem,
   objective, scope, or implementation need.
8. Keep content concise enough for presentation slides.

Validated digest:
{json.dumps(digest_data, ensure_ascii=False)}

Working solution and user angle:
{json.dumps(draft_data, ensure_ascii=False)}

Researched sources:
{json.dumps(source_summary, ensure_ascii=False)}

Return JSON only:
{json.dumps(schema, ensure_ascii=False)}
"""
    try:
        from groq import Groq

        model = os.environ.get("GROQ_MODEL", "openai/gpt-oss-120b")
        app.logger.info(
            "Groq storyline request started: model=%s sources=%d",
            model,
            len(source_summary),
        )
        completion = Groq(api_key=api_key, timeout=35, max_retries=1).chat.completions.create(
            model=model,
            messages=[
                {
                    "role": "system",
                    "content": "You are a senior proposal strategist. Return grounded JSON only.",
                },
                {"role": "user", "content": prompt},
            ],
            temperature=0,
            response_format={"type": "json_object"},
        )
        synthesized = json.loads(completion.choices[0].message.content)
        merged = deepcopy(fallback)
        for key in [
            "headline", "executive_summary", "problem_statement", "solution_summary",
            "why_us", "handed_over_artifact", "risk_note",
        ]:
            if isinstance(synthesized.get(key), str) and synthesized[key].strip():
                merged[key] = synthesized[key].strip()
        for key in [
            "pain_points", "approach_phases", "angle_application", "why_us_evidence",
            "market_insights",
        ]:
            if isinstance(synthesized.get(key), list) and synthesized[key]:
                merged[key] = synthesized[key]
        merged["case_objective"] = objective
        merged["quantified_target"] = quantified_target
        merged["why_now"] = why_now
        merged["scope_in"] = list(digest_data.get("scope_in") or [])
        merged["scope_out"] = list(digest_data.get("scope_out") or [])
        merged["client_angle"] = client_angle
        merged["market_context"] = source_summary
        known_sources = {
            source.get("title", "").strip().lower(): source
            for source in source_summary
            if source.get("title")
        }
        validated_insights = []
        for item in merged.get("market_insights") or []:
            if not isinstance(item, dict):
                continue
            source_title = str(item.get("source_title") or "").strip()
            source = known_sources.get(source_title.lower())
            insight = re.sub(r"\s+", " ", str(item.get("insight") or "")).strip()
            if not source or not insight:
                continue
            if len(insight) > 320:
                insight = insight[:317].rstrip() + "..."
            validated_insights.append(
                {
                    "insight": insight,
                    "source_title": source.get("title", source_title),
                    "url": source.get("url", ""),
                }
            )
            if len(validated_insights) >= 3:
                break
        merged["market_insights"] = validated_insights
        merged["storyline_model"] = model
        merged["storyline_request_id"] = getattr(completion, "id", "") or ""
        app.logger.info(
            "Groq storyline request completed: model=%s request_id=%s",
            model,
            merged["storyline_request_id"] or "unavailable",
        )
        return merged
    except Exception as error:
        app.logger.warning("Groq storyline synthesis failed; grounded fallback used (%s)", type(error).__name__)
        return fallback


@app.post("/api/storyline")
def storyline():
    payload = request.get_json() or {}
    digest_data = payload.get("digest") or {}
    if not str(digest_data.get("case_objective") or "").strip():
        return jsonify({"error": "A confirmed case objective is required before building the storyline."}), 422
    draft_data = payload.get("draft") or {}
    context_sources = (payload.get("context") or {}).get("sources", [])
    successful_sources = [source for source in context_sources if source.get("status") == "fetched"]
    return jsonify(synthesize_storyline(digest_data, draft_data, successful_sources))


@app.post("/api/render")
def render():
    payload = request.get_json() or {}
    digest_data = payload.get("digest") or {}
    if not str(digest_data.get("case_objective") or "").strip():
        return jsonify({"error": "A confirmed case objective is required before rendering the deck."}), 422
    draft_data = payload.get("draft") or {}
    story = payload.get("storyline") or {}
    incoming_brand = payload.get("brand_kit") or payload.get("brand_colors") or {}
    brand_kit = {
        **NEUTRAL_BRAND,
        **incoming_brand,
        "primary": valid_hex(incoming_brand.get("primary"), NEUTRAL_BRAND["primary"]),
        "secondary": valid_hex(incoming_brand.get("secondary"), NEUTRAL_BRAND["secondary"]),
        "accent": valid_hex(incoming_brand.get("accent"), NEUTRAL_BRAND["accent"]),
    }
    primary = brand_color(brand_kit["primary"])
    secondary = brand_color(brand_kit["secondary"])
    accent = brand_color(brand_kit["accent"])

    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    # Office uses the theme link color even when a run has an explicit color.
    theme_part = presentation.slide_master.part.part_related_by(RT.THEME)
    theme = parse_xml(theme_part.blob)
    for name in ("hlink", "folHlink"):
        for node in theme.xpath(f".//a:clrScheme/a:{name}"):
            for child in list(node):
                node.remove(child)
            color = readable_color(secondary, primary)
            node.append(parse_xml(
                f'<a:srgbClr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="{color}"/>'
            ))
    from lxml import etree
    theme_part._blob = etree.tostring(theme)
    blank = presentation.slide_layouts[6]
    title_slide = presentation.slides.add_slide(blank)
    add_background(title_slide, primary)
    add_logo(title_slide, brand_kit, Inches(10.5), Inches(0.45), Inches(1.8))
    add_text(title_slide, digest_data.get("client") or "Client proposal", Inches(0.9), Inches(0.8), Inches(9.4), Inches(0.5), accent, 16, True)
    add_text(title_slide, story.get("headline") or draft_data.get("quantified_headline") or "Case Builder Proposal", Inches(0.9), Inches(2.0), Inches(10.8), Inches(1.6), secondary, 32, True)
    add_text(
        title_slide,
        story.get("case_objective") or digest_data.get("case_objective") or story.get("executive_summary"),
        Inches(0.9),
        Inches(4.1),
        Inches(10.4),
        Inches(1.3),
        secondary,
        17,
    )
    bar = title_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(6.3), Inches(3.0), Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    summary_slide = presentation.slides.add_slide(blank)
    add_background(summary_slide, secondary)
    add_text(summary_slide, "Case summary", Inches(0.8), Inches(0.5), Inches(11.5), Inches(0.6), primary, 27, True)
    add_text(summary_slide, "Objective", Inches(0.8), Inches(1.35), Inches(2.0), Inches(0.35), accent, 14, True)
    add_text(
        summary_slide,
        story.get("case_objective") or digest_data.get("case_objective") or "Not explicitly stated in the source.",
        Inches(0.8),
        Inches(1.75),
        Inches(5.8),
        Inches(1.25),
        primary,
        17,
    )
    add_text(summary_slide, "Why act now", Inches(6.9), Inches(1.35), Inches(2.0), Inches(0.35), accent, 14, True)
    add_text(
        summary_slide,
        story.get("why_now") or digest_data.get("forcing_event") or "No explicit forcing event stated in the source.",
        Inches(6.9),
        Inches(1.75),
        Inches(5.2),
        Inches(1.25),
        primary,
        17,
    )
    add_text(summary_slide, "In scope", Inches(0.8), Inches(3.35), Inches(2.0), Inches(0.35), accent, 14, True)
    add_bullets(
        summary_slide,
        story.get("scope_in") or digest_data.get("scope_in") or [],
        Inches(0.8),
        Inches(3.8),
        Inches(5.8),
        Inches(2.5),
        primary,
        15,
    )
    target = story.get("quantified_target") or digest_data.get("quantified_target")
    add_text(summary_slide, "Success measure", Inches(6.9), Inches(3.35), Inches(2.2), Inches(0.35), accent, 14, True)
    add_text(
        summary_slide,
        target or "To be defined during alignment; no quantified target was stated.",
        Inches(6.9),
        Inches(3.8),
        Inches(5.2),
        Inches(1.2),
        primary,
        17,
    )

    problem_slide = presentation.slides.add_slide(blank)
    add_background(problem_slide, primary)
    add_text(problem_slide, "The problem to solve", Inches(0.8), Inches(0.5), Inches(11.4), Inches(0.6), secondary, 27, True)
    add_text(
        problem_slide,
        story.get("problem_statement") or draft_data.get("problem_statement"),
        Inches(0.8),
        Inches(1.35),
        Inches(11.4),
        Inches(1.0),
        accent,
        19,
        True,
    )
    add_text(problem_slide, "Key pain points", Inches(0.8), Inches(2.65), Inches(2.5), Inches(0.4), secondary, 15, True)
    add_bullets(
        problem_slide,
        story.get("pain_points") or digest_data.get("pain_points") or [],
        Inches(0.8),
        Inches(3.15),
        Inches(11.4),
        Inches(3.5),
        secondary,
        17,
    )

    solution_slide = presentation.slides.add_slide(blank)
    add_background(solution_slide, secondary)
    add_text(solution_slide, "Our proposed solution", Inches(0.8), Inches(0.5), Inches(11.4), Inches(0.6), primary, 27, True)
    add_text(
        solution_slide,
        story.get("solution_summary") or draft_data.get("solution_summary"),
        Inches(0.8),
        Inches(1.4),
        Inches(11.4),
        Inches(1.5),
        primary,
        18,
    )
    add_text(solution_slide, "How your angle changes the answer", Inches(0.8), Inches(3.2), Inches(5.0), Inches(0.4), accent, 15, True)
    angle_items = (
        story.get("angle_application")
        or draft_data.get("angle_application")
        or ([story.get("client_angle")] if story.get("client_angle") else [])
    )
    add_bullets(
        solution_slide,
        angle_items or ["No additional angle was supplied."],
        Inches(0.8),
        Inches(3.7),
        Inches(11.4),
        Inches(2.4),
        primary,
        16,
    )

    approach_slide = presentation.slides.add_slide(blank)
    add_background(approach_slide, primary)
    add_text(approach_slide, "Our approach", Inches(0.8), Inches(0.55), Inches(11.2), Inches(0.7), secondary, 27, True)
    phases = story.get("approach_phases") or draft_data.get("approach_phases") or []
    top = 1.55
    for index, phase in enumerate(phases[:5], 1):
        circle = approach_slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.85), Inches(top), Inches(0.55), Inches(0.55))
        circle.fill.solid()
        circle.fill.fore_color.rgb = accent
        circle.line.fill.background()
        add_text(approach_slide, str(index), Inches(1.02), Inches(top + 0.08), Inches(0.25), Inches(0.25), primary, 12, True)
        add_text(approach_slide, phase.get("name", ""), Inches(1.65), Inches(top - 0.02), Inches(2.7), Inches(0.65), secondary, 17, True)
        add_text(approach_slide, phase.get("duration", ""), Inches(4.55), Inches(top), Inches(1.4), Inches(0.35), accent, 12, True)
        add_text(approach_slide, phase.get("goal", ""), Inches(6.0), Inches(top - 0.02), Inches(6.2), Inches(0.65), secondary, 13)
        top += 1.0
    artifact = story.get("handed_over_artifact") or draft_data.get("handed_over_artifact")
    add_text(approach_slide, f"Client-owned output: {artifact}", Inches(0.85), Inches(6.75), Inches(11.5), Inches(0.45), secondary, 13, True)

    proof_slide = presentation.slides.add_slide(blank)
    add_background(proof_slide, secondary)
    add_text(proof_slide, "Why us", Inches(0.8), Inches(0.55), Inches(11.2), Inches(0.7), primary, 27, True)
    add_text(proof_slide, story.get("why_us") or draft_data.get("why_us"), Inches(0.8), Inches(1.5), Inches(11.4), Inches(1.5), primary, 17)
    evidence = story.get("why_us_evidence") or draft_data.get("evidence") or []
    add_text(proof_slide, "Grounded reasons", Inches(0.8), Inches(3.2), Inches(2.5), Inches(0.4), accent, 14, True)
    add_bullets(proof_slide, evidence, Inches(0.8), Inches(3.65), Inches(6.0), Inches(2.2), primary, 15)
    risk = story.get("risk_note") or draft_data.get("risk_note")
    risk_box = proof_slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(3.15), Inches(5.2), Inches(2.1))
    risk_box.fill.solid()
    risk_box.fill.fore_color.rgb = accent
    risk_box.line.fill.background()
    add_text(proof_slide, "Risk acknowledged", Inches(7.55), Inches(3.45), Inches(4.5), Inches(0.4), primary, 15, True)
    add_text(proof_slide, risk, Inches(7.55), Inches(4.0), Inches(4.35), Inches(0.95), primary, 13)

    research_slide = presentation.slides.add_slide(blank)
    add_background(research_slide, primary)
    add_text(research_slide, "Market evidence and implications", Inches(0.8), Inches(0.5), Inches(11.4), Inches(0.6), secondary, 27, True)
    insights = story.get("market_insights") or []
    top = 1.35
    for insight in insights[:4]:
        if isinstance(insight, dict):
            insight_text = insight.get("insight", "")
            source_label = insight.get("source_title", "")
        else:
            insight_text = str(insight)
            source_label = ""
        add_text(research_slide, f"• {insight_text}", Inches(0.8), Inches(top), Inches(11.4), Inches(0.60), secondary, 15)
        if source_label:
            citation = add_text(research_slide, f"Source: {source_label}", Inches(1.1), Inches(top + 0.70), Inches(10.8), Inches(0.40), accent, 12)
            if isinstance(insight, dict) and insight.get("url"):
                citation.text_frame.paragraphs[0].runs[0].hyperlink.address = insight["url"]
        top += 1.25
    if not insights:
        add_text(
            research_slide,
            "No external web evidence was available. The storyline remains grounded in the validated RFP and user input.",
            Inches(0.8),
            Inches(1.5),
            Inches(11.0),
            Inches(1.0),
            secondary,
            17,
        )
    cited_sources = [
        insight
        for insight in insights[:4]
        if isinstance(insight, dict) and insight.get("source_title")
    ]
    if cited_sources:
        add_text(research_slide, "Source titles above are clickable links.", Inches(0.8), Inches(6.75), Inches(11.4), Inches(0.35), secondary, 12)

    # Preserve excess content rather than clipping it or reducing it to tiny type.
    original_slides = list(presentation.slides)
    for source_slide in original_slides:
        remaining = "\n\n".join(getattr(source_slide.part, "_text_overflow", []))
        if not remaining:
            continue
        title = next((shape.text for shape in source_slide.shapes if shape.has_text_frame and shape.text), "Additional detail")
        while remaining:
            extra = presentation.slides.add_slide(blank)
            add_background(extra, primary)
            add_text(extra, "Additional detail", Inches(.8), Inches(.5), Inches(11.4), Inches(.6), secondary, 27, True)
            add_text(extra, title[:100], Inches(.8), Inches(1.2), Inches(11.4), Inches(.6), secondary, 16)
            visible, remaining, fitted_size = fit_text(remaining, Inches(11.4), Inches(4.9), 18)
            add_text(extra, visible, Inches(.8), Inches(2.0), Inches(11.4), Inches(4.9), secondary, fitted_size)

    output = io.BytesIO()
    presentation.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="proposal.pptx",
        mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )
def extract_file_content(uploaded_file):
    filename = (uploaded_file.filename or "").lower()
    file_bytes = uploaded_file.read()
    if not file_bytes:
        raise ValueError("The uploaded file is empty.")

    if filename.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(file_bytes))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        visual_text = (
            ""
            if find_brand_signals(text)
            else extract_pdf_visual_text(file_bytes)
        )
        return text, visual_text

    if filename.endswith(".pptx"):
        deck = Presentation(io.BytesIO(file_bytes))
        chunks = []
        for slide in deck.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    chunks.append(shape.text.strip())
        return "\n".join(chunks), ""

    if filename.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore"), ""

    raise ValueError("Upload a PDF, PPTX, or TXT file.")


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)))
