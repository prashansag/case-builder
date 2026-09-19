import io
import unittest

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

import app
from ppt_layout import fit_text, contrast_ratio, readable_color


class PptLayoutTests(unittest.TestCase):
    def test_low_contrast_brand_colors_are_corrected(self):
        for background in ("FFFFFF", "111111", "D46A3A"):
            color = RGBColor.from_string(background)
            self.assertGreaterEqual(contrast_ratio(readable_color(color, color), color), 4.5)

    def test_long_text_is_preserved_as_overflow(self):
        text = "A long objective requiring careful validation and implementation. " * 50
        visible, overflow, size = fit_text(text, Inches(3), Inches(.6), 18)
        self.assertTrue(overflow)
        self.assertGreaterEqual(size, 12)
        self.assertEqual(" ".join((visible + " " + overflow).split()), " ".join(text.split()))

    def test_deck_is_widescreen_and_all_text_stays_in_bounds(self):
        text = "Meaningful case detail about service delivery and patient access. " * 35
        digest = {"client": "Example", "case_objective": text, "pain_points": [text]}
        draft = app.build_draft(digest, app.STRUCTURE_LIBRARY[1])
        response = app.app.test_client().post("/api/render", json={
            "digest": digest, "draft": draft,
            "storyline": {"headline": text, "market_insights": [{
                "insight": text, "source_title": "Research",
                "url": "https://example.com/research",
            }]},
            "brand_kit": {"primary": "#FFFFFF", "secondary": "#FFFFFF", "accent": "#FFFFFF"},
        })
        self.assertEqual(response.status_code, 200)
        deck = Presentation(io.BytesIO(response.data))
        self.assertAlmostEqual(deck.slide_width / deck.slide_height, 16 / 9, places=5)
        self.assertGreater(len(deck.slides), 7)
        for slide in deck.slides:
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, 0)
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.left + shape.width, deck.slide_width)
                self.assertLessEqual(shape.top + shape.height, deck.slide_height)
                if shape.has_text_frame and shape.text:
                    paragraph = shape.text_frame.paragraphs[0]
                    size = paragraph.runs[0].font.size.pt
                    lines = shape.text.count("\n") + 1
                    self.assertLessEqual(lines * size * 1.25, shape.height.pt + .1)
                    self.assertGreaterEqual(
                        contrast_ratio(paragraph.runs[0].font.color.rgb, RGBColor(255, 255, 255)), 4.5
                    )