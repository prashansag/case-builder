import io
import unittest
from unittest.mock import patch

from pptx import Presentation

import app


class StorylineDeckTests(unittest.TestCase):
    def setUp(self):
        self.digest = {
            "client": "Northstar Health Network",
            "sector": "Healthcare",
            "geography": "Regional clinics",
            "case_objective": "Improve patient access and reduce appointment friction",
            "quantified_target": "Reduce average wait time by 20% within 12 months",
            "forcing_event": "Wait times are rising ahead of the new clinic launch.",
            "pain_points": [
                "Appointment delays are limiting patient access.",
                "Intake is fragmented across clinics.",
            ],
            "scope_in": ["Scheduling processes", "Clinic operations", "Digital intake"],
            "scope_out": [],
            "deliverables": ["Implementation roadmap", "Success measurement framework"],
        }

    def test_draft_is_grounded_in_case_not_reference_credentials(self):
        draft = app.build_draft(self.digest, app.STRUCTURE_LIBRARY[1])

        self.assertIn("Improve patient access", draft["case_objective"])
        self.assertIn("Appointment delays", draft["problem_statement"])
        self.assertEqual(draft["pain_points"], self.digest["pain_points"])
        self.assertEqual(draft["evidence"], [])
        self.assertNotIn("5,460", draft["why_us"])

    def test_enrichment_changes_solution_and_is_preserved(self):
        draft = app.build_draft(self.digest, app.STRUCTURE_LIBRARY[1])
        response = app.app.test_client().post(
            "/api/enrich",
            json={
                "draft": draft,
                "enrichment": (
                    "Lead with a patient-first service model. "
                    "Prioritize a reusable operating playbook."
                ),
            },
        )

        enriched = response.get_json()
        self.assertIn("patient-first", enriched["client_angle"])
        self.assertGreaterEqual(len(enriched["angle_application"]), 2)
        self.assertIn("design constraint", enriched["solution_summary"])

    def test_rendered_deck_contains_complete_storyline(self):
        draft = app.build_draft(self.digest, app.STRUCTURE_LIBRARY[1])
        with patch.dict("os.environ", {"GROQ_API_KEY": ""}, clear=False):
            story = app.synthesize_storyline(self.digest, draft, [])
        story["client_angle"] = "Lead with a patient-first service model."
        story["angle_application"] = ["Lead with a patient-first service model."]
        story["market_insights"] = [
            {
                "insight": "Digital intake can reduce avoidable booking friction.",
                "source_title": "Healthcare access research",
                "url": "https://example.com/research",
            }
        ]
        story["market_context"] = [
            {
                "title": "Healthcare access research",
                "url": "https://example.com/research",
            }
        ]

        response = app.app.test_client().post(
            "/api/render",
            json={
                "digest": self.digest,
                "draft": draft,
                "storyline": story,
                "brand_kit": {},
            },
        )

        self.assertEqual(response.status_code, 200)
        deck = Presentation(io.BytesIO(response.data))
        self.assertGreaterEqual(len(deck.slides), 7)
        all_text = "\n".join(
            shape.text
            for slide in deck.slides
            for shape in slide.shapes
            if hasattr(shape, "text")
        )
        for expected in [
            "Case summary",
            "Improve patient access",
            "Key pain points",
            "Appointment delays",
            "Our proposed solution",
            "patient-first",
            "Our approach",
            "Why us",
            "Market evidence and implications",
            "Digital intake",
        ]:
            self.assertIn(expected, all_text)

    def test_why_now_requires_a_real_trigger(self):
        generic = app.digest_text(
            "Competition is strong and the company wants growth. "
            "The objective is to improve customer experience.",
            "brief.txt",
        )
        explicit = app.digest_text(
            "The objective is to improve customer experience. "
            "The launch deadline is within 8 weeks, so the team must respond now.",
            "brief.txt",
        )

        self.assertEqual(generic["forcing_event"], "")
        self.assertIn("deadline", explicit["forcing_event"].lower())


if __name__ == "__main__":
    unittest.main()