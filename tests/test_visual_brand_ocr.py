import io
import unittest
from unittest.mock import patch

import pymupdf
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

import app


def uploaded(filename, content):
    return type(
        "Upload",
        (),
        {"filename": filename, "read": lambda self: content},
    )()


def image_only_brand_pdf(logo_text="Aashirvaad Atta"):
    font = ImageFont.truetype(
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        64,
    )
    logo = Image.new("RGB", (900, 180), "white")
    ImageDraw.Draw(logo).text(
        (28, 45),
        logo_text,
        font=font,
        fill="#9b241d",
    )
    logo_bytes = io.BytesIO()
    logo.save(logo_bytes, format="PNG")

    document = pymupdf.open()
    page = document.new_page(width=800, height=600)
    page.insert_text(
        (70, 300),
        "Request for proposal: consumer research and retail strategy",
        fontsize=16,
    )
    page.insert_image(
        pymupdf.Rect(70, 55, 700, 181),
        stream=logo_bytes.getvalue(),
    )
    content = document.tobytes()
    document.close()
    return content


class VisualBrandOcrTests(unittest.TestCase):
    def test_expanded_brand_catalog_maps_representative_aliases(self):
        representative_aliases = {
            "Nestlé India": "MAGGI",
            "Procter & Gamble": "Head and Shoulders",
            "PepsiCo": "LAY'S",
            "Coca-Cola": "CocaCola",
            "Mondelez India": "Cadbury DairyMilk",
            "Dabur India": "Hajmola",
            "Marico": "Saffola",
            "Godrej Consumer Products": "Good Knight",
            "Reckitt": "Dettol",
            "Colgate-Palmolive": "Colgate Toothpaste",
            "Gujarat Cooperative Milk Marketing Federation": "Amul Butter",
        }

        for expected_owner, logo_text in representative_aliases.items():
            with self.subTest(owner=expected_owner, logo_text=logo_text):
                signals = app.find_brand_signals(logo_text)
                self.assertEqual([signal["owner"] for signal in signals], [expected_owner])

    def test_common_ocr_spelling_variations_map_to_parent_owner(self):
        variations = {
            "Nestlé India": ["Kit Kat", "Nescafe"],
            "PepsiCo": ["Pepsi Co", "Lays Chips"],
            "Coca-Cola": ["Coca Cola", "Thumbs Up Cola"],
            "Mondelez India": ["Cadbury Dairymilk"],
            "Godrej Consumer Products": ["Goodknight"],
        }

        for expected_owner, aliases in variations.items():
            for alias in aliases:
                with self.subTest(owner=expected_owner, alias=alias):
                    signals = app.find_brand_signals(alias)
                    self.assertEqual([signal["owner"] for signal in signals], [expected_owner])

    def test_ambiguous_short_product_words_do_not_match(self):
        ambiguous_copy = (
            "Real results hit the target at high tide. The five star class "
            "will vanish after the expert review."
        )

        self.assertEqual(app.find_brand_signals(ambiguous_copy), [])

    def test_digest_endpoint_skips_image_only_logo_when_ocr_is_disabled(self):
        with patch.dict(
            "os.environ",
            {"ENABLE_PDF_OCR": "", "GROQ_API_KEY": ""},
            clear=False,
        ):
            response = app.app.test_client().post(
                "/api/digest",
                data={
                    "rfp_file": (
                        io.BytesIO(image_only_brand_pdf()),
                        "research-brief.pdf",
                    )
                },
                content_type="multipart/form-data",
            )

        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertEqual(payload["client"], "")
        self.assertFalse(payload["client_inferred_from_visual_evidence"])

    def test_known_text_brand_skips_pdf_ocr(self):
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text(
            (72, 72),
            "Prepared for Tata Salt. Scope includes retail research.",
            fontsize=14,
        )
        content = document.tobytes()
        document.close()

        with patch(
            "app.extract_pdf_visual_text",
            side_effect=AssertionError("OCR should not run"),
        ):
            text, visual_text = app.extract_file_content(
                uploaded("brief.pdf", content)
            )

        digest = app.digest_text(text, "brief.pdf", visual_text)
        self.assertEqual(visual_text, "")
        self.assertEqual(digest["client"], "Tata Consumer Products")
        self.assertFalse(digest["client_inferred_from_visual_evidence"])

    def test_filename_brand_does_not_hide_conflicting_image_logo(self):
        with patch(
            "app.extract_pdf_visual_text",
            return_value="Aashirvaad Atta",
        ) as extract_visual_text:
            text, visual_text = app.extract_file_content(
                uploaded("tata salt research brief.pdf", image_only_brand_pdf())
            )

        extract_visual_text.assert_called_once()
        digest = app.digest_text(
            text,
            "tata salt research brief.pdf",
            visual_text,
        )

        self.assertEqual(digest["client"], "ITC")
        self.assertEqual(digest["sector"], "FMCG")
        self.assertTrue(digest["client_inferred_from_visual_evidence"])
        self.assertIn("visual evidence", digest["client_evidence"][0].lower())
        self.assertNotIn("filename", digest["client_evidence"][0].lower())

    def test_filename_brand_is_labeled_as_fallback_evidence(self):
        digest = app.digest_text(
            "Request for proposal: consumer research and retail strategy",
            "tata salt research brief.pdf",
        )

        self.assertEqual(digest["client"], "Tata Consumer Products")
        self.assertFalse(digest["client_inferred_from_visual_evidence"])
        self.assertIn("source filename", digest["client_evidence"][0].lower())

    def test_text_and_pptx_extraction_remain_supported(self):
        text, visual_text = app.extract_file_content(
            uploaded("brief.txt", b"Client: Example Organisation")
        )
        self.assertEqual(text, "Client: Example Organisation")
        self.assertEqual(visual_text, "")

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(0, 0, 100, 100)
        box.text = "Client: Example Organisation"
        deck_bytes = io.BytesIO()
        presentation.save(deck_bytes)

        text, visual_text = app.extract_file_content(
            uploaded("brief.pptx", deck_bytes.getvalue())
        )
        self.assertIn("Client: Example Organisation", text)
        self.assertEqual(visual_text, "")


if __name__ == "__main__":
    unittest.main()